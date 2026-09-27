class DocumentExtractionError(Exception):
    pass


def extract_text(document):

    file_type = document.file_type

    if file_type == "txt":
        return extract_txt(document)

    if file_type == "pdf":
        return extract_pdf(document)

    if file_type == "docx":
        return extract_docx(document)

    if file_type == "xlsx":
        return extract_xlsx(document)

    if file_type in {"jpg", "jpeg", "png"}:
        raise DocumentExtractionError(
            "Image OCR will be enabled in Phase 2."
        )

    raise DocumentExtractionError(
        f"Unsupported file type: {file_type}"
    )


def extract_txt(document):

    document.file.seek(0)

    content = document.file.read()

    if isinstance(content, bytes):
        content = content.decode(
            "utf-8",
            errors="replace",
        )

    return str(content).strip()


def extract_pdf(document):

    try:
        import pypdf
    except ImportError:
        raise DocumentExtractionError(
            "pypdf is not installed."
        )

    document.file.seek(0)

    reader = pypdf.PdfReader(
        document.file
    )

    pages = []

    for page in reader.pages:

        text = page.extract_text() or ""

        if text.strip():
            pages.append(text.strip())

    result = "\n\n".join(pages).strip()

    if not result:
        raise DocumentExtractionError(
            "No readable text was found in this PDF."
        )

    return result


def extract_docx(document):

    try:
        from docx import Document as WordDocument
    except ImportError:
        raise DocumentExtractionError(
            "python-docx is not installed."
        )

    document.file.seek(0)

    word_document = WordDocument(
        document.file
    )

    paragraphs = []

    for paragraph in word_document.paragraphs:

        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    result = "\n\n".join(
        paragraphs
    ).strip()

    if not result:
        raise DocumentExtractionError(
            "No readable text was found in this DOCX."
        )

    return result


def extract_xlsx(document):

    try:
        from openpyxl import load_workbook
    except ImportError:
        raise DocumentExtractionError(
            "openpyxl is not installed."
        )

    document.file.seek(0)

    workbook = load_workbook(
        document.file,
        read_only=True,
        data_only=True,
    )

    sheets = []

    for worksheet in workbook.worksheets:

        rows = []

        for row in worksheet.iter_rows(
            values_only=True
        ):

            values = [
                str(value)
                for value in row
                if value is not None
            ]

            if values:
                rows.append(
                    " | ".join(values)
                )

        if rows:
            sheets.append(
                f"Sheet: {worksheet.title}\n"
                + "\n".join(rows)
            )

    workbook.close()

    result = "\n\n".join(
        sheets
    ).strip()

    if not result:
        raise DocumentExtractionError(
            "No readable data was found in this Excel file."
        )

    return result